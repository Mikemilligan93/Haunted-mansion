#!/usr/bin/env python3
"""
process_art.py - turns generated art sheets into game-ready sprites.

    python3 tools/process_art.py                 # process every sheet in tools/art_sheets.json
    python3 tools/process_art.py liam_td_walk    # process one sheet (by output name)
    python3 tools/process_art.py --preview       # also write 4x previews to tools/preview/
    python3 tools/process_art.py --palette-from-xlsx   # refresh tools/palette.json from the tracker

Pipeline per sheet:
  1. chroma-key the magenta (#FF00FF) background (plus pink fringe pixels)
  2. find the frames (fixed grid from art_sheets.json) and auto-crop / anchor them
  3. snap every pixel to the sheet's palette group (tools/palette.json)
  4. downscale with a block filter: each output pixel takes the most common palette colour
     of its source block (nearest-neighbour-style, crisp; dark outlines get priority so lines survive)
  5. tiles: trim the cell edges and check that tiles repeat without seams
Output: assets/sprites/<name>.png (frames packed in a grid) + assets/sprites/manifest.json

Requires Python 3 with Pillow and numpy (pip3 install pillow numpy). openpyxl only for --palette-from-xlsx.
"""
import json
import sys
from pathlib import Path

import numpy as np
from PIL import Image

ROOT = Path(__file__).resolve().parent.parent
SRC_DIR = ROOT / "assets" / "img"
OUT_DIR = ROOT / "assets" / "sprites"
TOOLS = ROOT / "tools"
PREVIEW_DIR = TOOLS / "preview"


# ------------------------------------------------------------------ palette
def hex_rgb(h):
    h = h.lstrip("#")
    return tuple(int(h[i:i + 2], 16) for i in (0, 2, 4))


def rgb_to_lab(rgb):
    """rgb: (..., 3) uint8 -> (..., 3) float Lab (D65)."""
    c = rgb.astype(np.float64) / 255.0
    c = np.where(c > 0.04045, ((c + 0.055) / 1.055) ** 2.4, c / 12.92)
    m = np.array([[0.4124, 0.3576, 0.1805], [0.2126, 0.7152, 0.0722], [0.0193, 0.1192, 0.9505]])
    xyz = c @ m.T / np.array([0.95047, 1.0, 1.08883])
    f = np.where(xyz > 0.008856, np.cbrt(xyz), 7.787 * xyz + 16 / 116)
    return np.stack([116 * f[..., 1] - 16, 500 * (f[..., 0] - f[..., 1]), 200 * (f[..., 1] - f[..., 2])], -1)


class Palette:
    def __init__(self, colours):
        self.rgb = np.array([hex_rgb(c["hex"]) for c in colours], dtype=np.uint8)
        self.lab = rgb_to_lab(self.rgb)
        lum = self.lab[:, 0]
        self.ink = int(np.argmin(lum))  # darkest colour = outline ink
        self.dark = np.nonzero(lum < 22)[0]  # ink-like colours (outline, eye rings, deepest shadows)

    def index(self, rgb):
        """rgb (N,3) -> nearest palette index (N,)"""
        lab = rgb_to_lab(rgb)
        d = ((lab[:, None, :] - self.lab[None, :, :]) ** 2).sum(-1)
        return d.argmin(1)


def load_palettes():
    data = json.loads((TOOLS / "palette.json").read_text())
    return {k: Palette(v) for k, v in data["groups"].items()}


# ------------------------------------------------------------------ chroma key
def key_mask(arr):
    """True where the pixel is art (not magenta background or pink fringe)."""
    r, g, b = (arr[..., i].astype(np.int32) for i in range(3))
    near = (255 - r) ** 2 + g ** 2 + (255 - b) ** 2 < 120 ** 2
    fringe = (r > 90) & (b > 90) & (g < 0.6 * np.minimum(r, b)) & (np.abs(r - b) < 110)
    return ~(near | fringe)


# ------------------------------------------------------------------ block downscale
def downscale(idx, alpha, out_w, out_h, x0, y0, sx, sy, n_colours, ink, ink_share, alpha_min, dark=None):
    """
    idx: (H,W) palette indices of the source; alpha: (H,W) bool.
    Output pixel (u,v) covers source rect [x0+u*sx, x0+(u+1)*sx) x [y0+v*sy, ...).
    Returns (out_h,out_w) index array with -1 = transparent.
    """
    H, W = idx.shape
    out = np.full((out_h, out_w), -1, dtype=np.int32)
    for v in range(out_h):
        ya, yb = int(round(y0 + v * sy)), int(round(y0 + (v + 1) * sy))
        ya, yb = max(ya, 0), min(max(yb, ya + 1), H)
        if ya >= H or yb <= 0:
            continue
        for u in range(out_w):
            xa, xb = int(round(x0 + u * sx)), int(round(x0 + (u + 1) * sx))
            xa, xb = max(xa, 0), min(max(xb, xa + 1), W)
            if xa >= W or xb <= 0:
                continue
            a = alpha[ya:yb, xa:xb]
            if a.size == 0 or a.mean() < alpha_min:
                continue
            cols = idx[ya:yb, xa:xb][a]
            counts = np.bincount(cols, minlength=n_colours)
            if ink is not None and dark is not None and len(dark) and counts[dark].sum() >= ink_share * a.size:
                out[v, u] = int(dark[counts[dark].argmax()])
            elif ink is not None and counts[ink] >= ink_share * a.size:
                out[v, u] = ink
            else:
                out[v, u] = int(counts.argmax())
    return out


def to_rgba(out_idx, pal):
    h, w = out_idx.shape
    img = np.zeros((h, w, 4), dtype=np.uint8)
    m = out_idx >= 0
    img[m, :3] = pal.rgb[out_idx[m]]
    img[m, 3] = 255
    return img


def drop_specks(img, min_px):
    """Remove small detached pixel islands (bleed from neighbouring cells, stray noise). Keeps the largest island."""
    a = img[..., 3] > 0
    h, w = a.shape
    label = np.zeros((h, w), dtype=np.int32)
    sizes = [0]
    for y in range(h):
        for x in range(w):
            if a[y, x] and not label[y, x]:
                n = len(sizes)
                stack, size = [(y, x)], 0
                label[y, x] = n
                while stack:
                    cy, cx = stack.pop()
                    size += 1
                    for ny, nx in ((cy + 1, cx), (cy - 1, cx), (cy, cx + 1), (cy, cx - 1), (cy + 1, cx + 1), (cy - 1, cx - 1), (cy + 1, cx - 1), (cy - 1, cx + 1)):
                        if 0 <= ny < h and 0 <= nx < w and a[ny, nx] and not label[ny, nx]:
                            label[ny, nx] = n
                            stack.append((ny, nx))
                sizes.append(size)
    if len(sizes) <= 2:
        return img
    biggest = int(np.argmax(sizes))
    out = img.copy()
    for n, size in enumerate(sizes):
        if n and n != biggest and size < min_px:
            out[label == n] = 0
    return out


def add_outline(img, ink_rgb):
    """Paint a 1px ink outline on silhouette edge pixels (inside the sprite)."""
    a = img[..., 3] > 0
    pad = np.pad(a, 1)
    edge = a & ~(pad[:-2, 1:-1] & pad[2:, 1:-1] & pad[1:-1, :-2] & pad[1:-1, 2:])
    out = img.copy()
    out[edge, :3] = ink_rgb
    return out


# ------------------------------------------------------------------ sheet kinds
def cells(sheet, w, h):
    rows, cols = sheet["grid"]
    cw, ch = w / cols, h / rows
    for r in range(rows):
        for c in range(cols):
            yield r, c, int(round(c * cw)), int(round(r * ch)), int(round(cw)), int(round(ch))


def bbox(mask):
    ys, xs = np.nonzero(mask)
    if len(xs) == 0:
        return None
    return xs.min(), ys.min(), xs.max() + 1, ys.max() + 1


def cell_masks(mask, cell_list):
    """
    Which art pixels belong to which frame. Generated sheets often let a drawing poke over its
    cell line (hair into the row above, a cape into the next column), so instead of cutting at
    the grid we keep every connected blob whole and give it to the cell that holds most of it.
    Needs scipy for speed; without it, falls back to cutting at the cell lines.
    """
    out = {}
    try:
        from scipy import ndimage
    except ImportError:
        print("     (scipy not installed: cutting frames at the cell lines; pip3 install scipy for better results)")
        for r, c, x, y, cw, ch in cell_list:
            m = np.zeros_like(mask)
            m[y:y + ch, x:x + cw] = mask[y:y + ch, x:x + cw]
            out[(r, c)] = m
        return out
    labels, n = ndimage.label(mask, structure=np.ones((3, 3)))
    owner = np.full(n + 1, -1)
    overlap = np.zeros((len(cell_list), n + 1), dtype=np.int64)
    for i, (r, c, x, y, cw, ch) in enumerate(cell_list):
        overlap[i] = np.bincount(labels[y:y + ch, x:x + cw].ravel(), minlength=n + 1)
    owner[1:] = overlap[:, 1:].argmax(0)
    owner_map = owner[labels]
    for i, (r, c, x, y, cw, ch) in enumerate(cell_list):
        out[(r, c)] = mask & (owner_map == i)
    return out


def process_sprite(sheet, arr, idx, mask, pal, scale_cache):
    """Characters: anchored at the cell's centre x and the row's baseline (feet)."""
    H, W = mask.shape
    rows, cols = sheet["grid"]
    fw, fh = sheet["frame"]
    anchor = sheet.get("anchor", "feet")
    cell_list = list(cells(sheet, W, H))
    owned = cell_masks(mask, cell_list)
    # per-cell bounding boxes (in sheet coordinates), from the blobs each cell owns
    boxes = {(r, c): bbox(owned[(r, c)]) for r, c, *_ in cell_list}

    # scale
    if "scale_from" in sheet:
        s = scale_cache[sheet["scale_from"]]
    else:
        rr, cc = sheet["scale_ref"]["frame"]
        b = boxes[(rr, cc)]
        dim = sheet["scale_ref"].get("dim", "height")
        size = (b[3] - b[1]) if dim == "height" else (b[2] - b[0])
        s = sheet["scale_ref"]["px"] / size
    scale_cache[sheet["name"]] = s

    out = np.zeros((rows * fh, cols * fw, 4), dtype=np.uint8)
    bounds = {}
    ax, ay = fw / 2, fh - 2
    for r in range(rows):
        row_boxes = [boxes[(r, c)] for c in range(cols) if boxes[(r, c)]]
        if not row_boxes:
            continue
        if anchor == "feet":
            base = np.median([b[3] for b in row_boxes]) if sheet.get("baseline") == "median" else max(b[3] for b in row_boxes)
        else:  # centre
            base = None
        for c in range(cols):
            _, _, x, y, cw, ch = next(t for t in cell_list if t[0] == r and t[1] == c)
            cx = x + cw / 2 + sheet.get("x_offset", 0)
            if anchor == "feet":
                x0, y0 = cx - ax / s, base - ay / s
            else:
                cy = y + ch / 2
                x0, y0 = cx - (fw / 2) / s, cy - (fh / 2) / s
            # only this frame's own blobs, so neighbours never bleed in and nothing is cut off
            o = downscale(idx, owned[(r, c)], fw, fh, x0, y0, 1 / s, 1 / s, len(pal.rgb), pal.ink,
                          sheet.get("ink_share", 0.25), sheet.get("alpha_min", 0.5), pal.dark)
            img = drop_specks(to_rgba(o, pal), sheet.get("min_blob", 10))
            if sheet.get("outline", True):
                img = add_outline(img, pal.rgb[pal.ink])
            out[r * fh:(r + 1) * fh, c * fw:(c + 1) * fw] = img
            b = bbox(img[..., 3] > 0)
            bounds.setdefault(r, {})[c] = None if b is None else [int(v) for v in b]
    # bounds[row][col] = [x0, y0, x1, y1] of the drawing inside its frame (used to line up animations)
    return out, {"scale": round(s, 4), "bounds": [[bounds.get(r, {}).get(c) for c in range(cols)] for r in range(rows)]}


def process_fit(sheet, arr, idx, mask, pal):
    """Props: union bbox of all frames, cropped identically from every cell, fitted to a target size."""
    H, W = mask.shape
    rows, cols = sheet["grid"]
    cl = list(cells(sheet, W, H))
    owned = cell_masks(mask, cl)
    ub = None
    for r, c, x, y, cw, ch in cl:
        bb = bbox(owned[(r, c)])
        if bb is None:
            continue
        b = (bb[0] - x, bb[1] - y, bb[2] - x, bb[3] - y)   # relative to the cell
        ub = b if ub is None else (min(ub[0], b[0]), min(ub[1], b[1]), max(ub[2], b[2]), max(ub[3], b[3]))
    bw, bh = ub[2] - ub[0], ub[3] - ub[1]
    tw, th = sheet["fit"]
    if tw and th:
        fw, fh = tw, th
    elif tw:
        fw, fh = tw, int(round(bh * tw / bw))
    else:
        fh, fw = th, int(round(bw * th / bh))
    sx, sy = bw / fw, bh / fh
    out = np.zeros((rows * fh, cols * fw, 4), dtype=np.uint8)
    for r, c, x, y, cw, ch in cl:
        o = downscale(idx, owned[(r, c)], fw, fh, x + ub[0], y + ub[1], sx, sy, len(pal.rgb), pal.ink,
                      sheet.get("ink_share", 0.35), sheet.get("alpha_min", 0.5))
        out[r * fh:(r + 1) * fh, c * fw:(c + 1) * fw] = drop_specks(to_rgba(o, pal), sheet.get("min_blob", 6))
    return out, {"frame": [fw, fh]}


def process_cellcrop(sheet, arr, idx, mask, pal):
    """Portraits: the same rectangle cut from every cell, scaled to a fixed size."""
    H, W = mask.shape
    rows, cols = sheet["grid"]
    rx, ry, rw, rh = sheet["crop"]
    fw, fh = sheet["size"]
    out = np.zeros((rows * fh, cols * fw, 4), dtype=np.uint8)
    cl = list(cells(sheet, W, H))
    owned = cell_masks(mask, cl)
    for r, c, x, y, cw, ch in cl:
        o = downscale(idx, owned[(r, c)], fw, fh, x + rx, y + ry, rw / fw, rh / fh, len(pal.rgb), pal.ink,
                      sheet.get("ink_share", 0.25), 0.5, pal.dark)
        img = to_rgba(o, pal)
        if sheet.get("outline", False):
            img = add_outline(img, pal.rgb[pal.ink])
        out[r * fh:(r + 1) * fh, c * fw:(c + 1) * fw] = img
    return out, {"frame": [fw, fh]}


def process_tiles(sheet, arr, idx, mask, pal):
    """Opaque tiles from a fixed grid. Cell edges are trimmed (removes drawn seams) and the result is checked for seams."""
    H, W = mask.shape
    tw, th = sheet["tile"]
    inset = sheet.get("inset", 0)
    regions = sheet["regions"]
    frames, report = [], []
    for reg in regions:
        rx, ry, rw, rh = reg["rect"]
        n = reg["count"]
        cw = rw / n
        for i in range(n):
            x0 = rx + i * cw + inset
            y0 = ry + inset
            sx = (cw - 2 * inset) / reg["size"][0]
            sy = (rh - 2 * inset) / reg["size"][1]
            full = np.ones_like(mask)
            o = downscale(idx, full, reg["size"][0], reg["size"][1], x0, y0, sx, sy, len(pal.rgb), None, 1, 0.0)
            frames.append((reg["name"], i, to_rgba(o, pal)))
    # pack: one row per region, padded to the widest
    widths = {}
    for name, i, img in frames:
        widths.setdefault(name, []).append(img)
    rows = list(widths.items())
    fw = max(r[1][0].shape[1] for r in rows)
    fh = max(r[1][0].shape[0] for r in rows)
    ncol = max(len(r[1]) for r in rows)
    out = np.zeros((len(rows) * fh, ncol * fw, 4), dtype=np.uint8)
    layout = {}
    for ri, (name, imgs) in enumerate(rows):
        layout[name] = {"row": ri, "count": len(imgs), "size": [imgs[0].shape[1], imgs[0].shape[0]]}
        for ci, img in enumerate(imgs):
            out[ri * fh:ri * fh + img.shape[0], ci * fw:ci * fw + img.shape[1]] = img
            if sheet.get("seamless_check") and f"{name}[{ci}]" not in sheet.get("edge_tiles", []):
                lum = img[..., :3].mean(-1)
                if img.shape[1] > 2:
                    jump = abs(lum[:, 0].mean() - lum[:, 1:-1].mean()) + abs(lum[:, -1].mean() - lum[:, 1:-1].mean())
                    if jump > 18:
                        report.append(f"{name}[{ci}] edge columns differ from the inside by {jump:.0f} (possible seam)")
    return out, {"frame": [fw, fh], "regions": layout, "seam_report": report}


# ------------------------------------------------------------------ main
def palette_from_xlsx():
    from openpyxl import load_workbook
    wb = load_workbook(ROOT / "docs" / "Haunted_Mansion_Art_Prompts.xlsx", read_only=True)
    groups = {}
    for row in wb["Palette"].iter_rows(min_row=2, values_only=True):
        name, hx, _, who = row[:4]
        if not hx or hx.upper() == "#FF00FF":
            continue
        groups.setdefault((who or "other").lower(), []).append({"name": name, "hex": hx.upper()})
    outline = {"name": "Outline and hatching", "hex": "#1A1020"}
    shared = groups.get("shared", [])
    brass = [c for c in groups.get("ambrose", []) if "brass" in c["name"] or "wax" in c["name"]]
    pal = {"_note": "Generated from the Palette sheet of docs/Haunted_Mansion_Art_Prompts.xlsx. Regenerate with: python3 tools/process_art.py --palette-from-xlsx",
           "groups": {"liam": groups.get("liam", []),
                      "ambrose": [outline] + groups.get("ambrose", []) + shared,
                      "world": [outline] + groups.get("world", []) + brass + shared}}
    (TOOLS / "palette.json").write_text(json.dumps(pal, indent=1))
    print("palette.json refreshed:", {k: len(v) for k, v in pal["groups"].items()})


def main(argv):
    if "--palette-from-xlsx" in argv:
        palette_from_xlsx()
        return
    preview = "--preview" in argv
    only = [a for a in argv if not a.startswith("--")]
    sheets = json.loads((TOOLS / "art_sheets.json").read_text())["sheets"]
    pals = load_palettes()
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    man_path = OUT_DIR / "manifest.json"
    manifest = json.loads(man_path.read_text()) if man_path.exists() else {"sheets": {}}
    scale_cache = {k: v.get("scale") for k, v in manifest["sheets"].items() if v.get("scale")}
    problems = []
    for sh in sheets:
        if only and sh["name"] not in only:
            continue
        src = SRC_DIR / sh["file"]
        if not src.exists():
            problems.append(f"{sh['name']}: missing source {sh['file']} (the game will use fallback art)")
            continue
        im = Image.open(src).convert("RGB")
        arr = np.asarray(im)
        mask = key_mask(arr)
        if sh["kind"] == "tiles":
            mask[:] = True
        pal = pals[sh["palette"]]
        if sh.get("extra_colors"):
            # a sheet can add a few colours to its palette group (e.g. white china on a cream cloth)
            base = json.loads((TOOLS / "palette.json").read_text())["groups"][sh["palette"]]
            pal = Palette(base + [{"name": "extra", "hex": h} for h in sh["extra_colors"]])
        flat = arr.reshape(-1, 3)
        uniq, inv = np.unique(flat, axis=0, return_inverse=True)
        idx = pal.index(uniq)[inv.ravel()].reshape(mask.shape)
        if sh["kind"] == "sprite":
            out, meta = process_sprite(sh, arr, idx, mask, pal, scale_cache)
            meta["frame"] = sh["frame"]
            meta["origin"] = [sh["frame"][0] // 2, sh["frame"][1] - 2] if sh.get("anchor", "feet") == "feet" else [sh["frame"][0] // 2, sh["frame"][1] // 2]
        elif sh["kind"] == "fit":
            out, meta = process_fit(sh, arr, idx, mask, pal)
            meta["origin"] = [meta["frame"][0] // 2, meta["frame"][1]]
        elif sh["kind"] == "cellcrop":
            out, meta = process_cellcrop(sh, arr, idx, mask, pal)
            meta["origin"] = [0, 0]
        else:
            out, meta = process_tiles(sh, arr, idx, mask, pal)
            problems += [f"{sh['name']}: {m}" for m in meta.pop("seam_report")]
        if sh.get("remap"):
            # swap palette colours after processing (e.g. cream china -> white so it reads on a cream cloth)
            for a, b in sh["remap"].items():
                ar, br = hex_rgb(a), hex_rgb(b)
                hit = (out[..., 3] > 0) & (out[..., 0] == ar[0]) & (out[..., 1] == ar[1]) & (out[..., 2] == ar[2])
                out[hit, :3] = br
        Image.fromarray(out, "RGBA").save(OUT_DIR / f"{sh['name']}.png")
        entry = {"image": f"{sh['name']}.png", "grid": sh.get("grid"), **meta}
        for k in ("anims", "rows", "names", "skip"):
            if k in sh:
                entry[k] = sh[k]
        manifest["sheets"][sh["name"]] = entry
        if preview:
            PREVIEW_DIR.mkdir(exist_ok=True)
            p = Image.fromarray(out, "RGBA")
            bg = Image.new("RGBA", p.size, (60, 60, 70, 255))
            bg.alpha_composite(p)
            bg.resize((p.width * 4, p.height * 4), Image.NEAREST).save(PREVIEW_DIR / f"{sh['name']}.png")
        print(f"ok  {sh['name']:<26} {out.shape[1]}x{out.shape[0]}  {meta.get('scale', '')}")
    man_path.write_text(json.dumps(manifest, indent=1))
    for p in problems:
        print("WARN", p)


if __name__ == "__main__":
    main(sys.argv[1:])
